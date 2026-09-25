from pypdf import PdfReader
from docx import Document as doc
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from constants.constants import root_path
import os

async def load_files(subject):
    print("Loading Files")
    path=os.path.join(root_path,subject)
    files=[]
    try :
        folder=os.scandir(path)
        for file in folder:
            files.append(file.path)
    except:
        return {"Wrong file path"}
        
    try:
        loaded_files=[]
        
        for file in files:
            content=[]
            
            if file.endswith('pdf'):
                pdf=PdfReader(file)
                for page in pdf.pages:
                    content.append(Document(page_content=page.extract_text(),metadata={"source":file,"subject":subject}))
                loaded_files.extend(content)
            
            elif file.endswith('docx'):
                f= open(file,'rb')
                document=doc(f)
                for page in document.paragraphs:
                    content.append(Document(page_content=page.text,metadata={"source":file,"subject":subject}))
                loaded_files.extend(content)
            
            elif file.endswith('txt') or file.endswith('md'):
                with open(file,'r',encoding='utf-8') as f:
                    content=Document(page_content=f.read(),metadata={"source":file,"subject":subject})
                    loaded_files.extend(content)

            else:
                continue
            
        return loaded_files
    except:
        return {"Files not found"}
    

async def chunking_data(documents):
    print("Chunking data")
    try:
        splitter=RecursiveCharacterTextSplitter(chunk_size=1000,chunk_overlap=200)
        chunk=splitter.split_documents(documents)
        return chunk
    except:
        return {"Data not found"}

